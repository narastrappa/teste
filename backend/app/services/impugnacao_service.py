"""Impugnação service: IA analysis, geração de minuta DOCX."""
import logging
from datetime import datetime, timezone
from typing import Any, Dict
from uuid import UUID

import anthropic
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.config import settings
from app.models.impugnacao import Impugnacao, StatusImpugnacaoEnum
from app.models.edital import Edital

logger = logging.getLogger(__name__)


async def analisar_impugnacao(db: AsyncSession, impugnacao: Impugnacao) -> Dict[str, Any]:
    """Análise jurídica da impugnação via IA."""
    result = await db.execute(select(Edital).where(Edital.id == impugnacao.edital_id))
    edital = result.scalar_one_or_none()

    prompt = f"""Você é um advogado especializado em licitações públicas brasileiras (Lei 14.133/2021 e Lei 8.666/93).

Analise os motivos de impugnação abaixo para o edital e forneça:
1. Análise jurídica de cada motivo
2. Fundamentos legais aplicáveis
3. Chances de êxito (Alta/Média/Baixa)
4. Recomendações para fortalecer a impugnação
5. Pontos críticos a observar

EDITAL:
- Número: {edital.numero if edital else 'N/A'}
- Órgão: {edital.orgao if edital else 'N/A'}
- Objeto: {edital.objeto if edital else 'N/A'}
- Modalidade: {edital.modalidade.value if edital else 'N/A'}

MOTIVOS DA IMPUGNAÇÃO:
{impugnacao.motivos}

Responda em JSON:
{{
  "analise_juridica": "análise detalhada",
  "fundamentos_legais": ["art. X da Lei Y", ...],
  "chance_exito": "Alta|Média|Baixa",
  "recomendacoes": ["..."],
  "pontos_criticos": ["..."],
  "argumentos_contrarios": ["possíveis contra-argumentos do órgão"]
}}"""

    client = anthropic.Anthropic(api_key=settings.anthropic_api_key)

    try:
        response = client.messages.create(
            model=settings.anthropic_model,
            max_tokens=2000,
            messages=[{"role": "user", "content": prompt}],
        )
        import json
        content = response.content[0].text
        start = content.find("{")
        end = content.rfind("}") + 1
        analise = json.loads(content[start:end])
    except Exception as e:
        logger.error(f"Erro na análise IA da impugnação {impugnacao.id}: {e}")
        analise = {
            "analise_juridica": "Análise indisponível",
            "fundamentos_legais": [],
            "chance_exito": "Média",
            "recomendacoes": [],
            "pontos_criticos": [str(e)],
            "argumentos_contrarios": [],
        }

    return {"analise": analise, "modelo": settings.anthropic_model}


async def gerar_minuta_docx(db: AsyncSession, impugnacao: Impugnacao) -> str:
    """Gera minuta de impugnação em DOCX e retorna s3_key."""
    from docx import Document
    from docx.shared import Pt, Inches
    import io
    from app.storage import storage_client

    result = await db.execute(select(Edital).where(Edital.id == impugnacao.edital_id))
    edital = result.scalar_one_or_none()

    doc = Document()

    # Title
    title = doc.add_heading("IMPUGNAÇÃO AO EDITAL", 0)
    title.alignment = 1  # CENTER

    doc.add_paragraph()

    # Header info
    doc.add_paragraph(f"Edital nº: {edital.numero if edital else 'N/A'}")
    doc.add_paragraph(f"Órgão: {edital.orgao if edital else 'N/A'}")
    doc.add_paragraph(f"Modalidade: {edital.modalidade.value if edital else 'N/A'}")
    doc.add_paragraph(f"Data de abertura: {edital.data_abertura.strftime('%d/%m/%Y às %H:%M') if edital and edital.data_abertura else 'N/A'}")
    doc.add_paragraph(f"Prazo para impugnação: {impugnacao.prazo_impugnacao.strftime('%d/%m/%Y') if impugnacao.prazo_impugnacao else 'N/A'}")

    doc.add_paragraph()
    doc.add_heading("EXMO. SR. PREGOEIRO / COMISSÃO DE LICITAÇÃO", 2)
    doc.add_paragraph()

    # Opening
    intro = doc.add_paragraph(
        f"Vem respeitosamente perante V.Sa., a empresa EMPRESA LICITANTE, por seus procuradores, com fundamento no art. 164 da Lei nº 14.133/2021, "
        f"interpor a presente IMPUGNAÇÃO ao Edital nº {edital.numero if edital else 'N/A'}, "
        f"pelas razões de fato e de direito a seguir expostas:"
    )

    doc.add_paragraph()
    doc.add_heading("I - DOS FATOS", 1)

    # Motivos
    doc.add_paragraph(impugnacao.motivos)

    # IA analysis
    if impugnacao.analise_ia:
        analise = impugnacao.analise_ia.get("analise", {})

        doc.add_paragraph()
        doc.add_heading("II - DO DIREITO", 1)

        if analise.get("analise_juridica"):
            doc.add_paragraph(analise["analise_juridica"])

        if analise.get("fundamentos_legais"):
            doc.add_paragraph()
            doc.add_heading("Fundamentos Legais:", 3)
            for fund in analise["fundamentos_legais"]:
                doc.add_paragraph(f"• {fund}", style="List Bullet")

        doc.add_paragraph()
        doc.add_heading("III - DOS PEDIDOS", 1)
        doc.add_paragraph(
            "Diante do exposto, requer-se:"
            "\na) O conhecimento e o provimento da presente impugnação;"
            "\nb) A retificação/suspensão do edital nos pontos impugnados;"
            "\nc) A reabertura do prazo para apresentação de propostas, se necessário."
        )

    doc.add_paragraph()
    doc.add_paragraph(f"Local e Data: _____________________, {datetime.now().strftime('%d de %B de %Y')}")
    doc.add_paragraph()
    doc.add_paragraph("_______________________________")
    doc.add_paragraph("Responsável Legal / Procurador")

    # Save to bytes
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    s3_key = f"impugnacoes/{impugnacao.id}/minuta_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    storage_client.upload(
        s3_key,
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    return s3_key
